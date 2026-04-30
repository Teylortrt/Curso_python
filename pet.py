from openpyxl import Workbook, load_workbook
import os

arquivo = "petshop.xlsx"

if not os.path.exists(arquivo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Atendimentos"
    
    
    ws.append(["Nome do Cachorro", "Serviço"])
    
    wb.save(arquivo)

wb = load_workbook(arquivo) # vai carrega o arquivo
ws = wb["Atendimentos"]

while True:
    print("\n--- Cadastro Petshop ---")
    
    nome = input("Nome do cachorro: ")
    servico = input("Tipo de serviço (Banho/Tosa): ")
    
    
    ws.append([nome, servico])
    

    wb.save(arquivo)
    
    print("✅ Registro salvo!")
    while True:
        try:   
            continuar = input("Deseja cadastrar outro pet? (s/n): ").lower()#lower para converter a resposta para minúscula
            
            break
        except ValueError:
            print("Por favor, responda com 's' para sim ou 'n' para não.")

print("Dados salvos no arquivo petshop.xlsx")